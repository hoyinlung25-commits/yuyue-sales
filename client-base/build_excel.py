#!/usr/bin/env python3
"""Build client_base.xlsx — policy book screenshot data only."""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).parent
OUTPUT = BASE / "client_base.xlsx"

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14)

# Bilingual headers matching your screenshots
POLICY_HEADERS = [
    "銷售團隊名稱\nSales Team",
    "代理人\nAgent",
    "權益人\nPolicyowner",
    "受保人\nInsured",
    "保單號碼\nPolicy Number",
    "基本計劃\nBasic Plan",
    "投保日期\nApplication Date",
    "簽發日期\nIssue Date",
    "保單狀態\nPolicy Status",
    "供款年期\nPremium Term (Years)",
    "保單到期日\nPolicy Expiry Date",
    "運單號碼\nTracking Number",
]


def read_csv(name: str) -> tuple[list[str], list[list]]:
    path = BASE / name
    with path.open(encoding="utf-8") as f:
        rows = list(csv.reader(f))
    return rows[0], rows[1:]


def style_sheet(ws, ncols: int, nrows: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{max(nrows, 1)}"
    for col in range(1, ncols + 1):
        letter = get_column_letter(col)
        max_len = 14
        for row in range(1, min(nrows + 1, 200)):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                max_len = max(max_len, min(len(str(val)) + 2, 55))
        ws.column_dimensions[letter].width = max_len
    ws.row_dimensions[1].height = 36


def write_table(ws, headers: list[str], data: list[list], start_row: int = 1) -> int:
    for c, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=c, value=h)
    for r, row in enumerate(data, start_row + 1):
        for c, val in enumerate(row, 1):
            ws.cell(row=r, column=c, value=val if val != "" else None)
    return start_row + len(data)


def build_policies(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Policies"
    _, data = read_csv("policies_register.csv")
    nrows = write_table(ws, POLICY_HEADERS, data)
    style_sheet(ws, len(POLICY_HEADERS), nrows)


def build_clients_summary(wb: Workbook) -> None:
    """One row per unique policyowner — counts only, from screenshot data."""
    ws = wb.create_sheet("Clients Summary")
    headers = [
        "權益人\nPolicyowner",
        "受保人\nInsured",
        "Inforce Policies",
        "Plans (from screenshots)",
    ]
    _, rows = read_csv("policies_register.csv")
    # columns: policyowner idx 2, insured 3, status 8, plan 5
    clients: dict[str, dict] = {}
    for row in rows:
        owner = row[2]
        if not owner:
            continue
        if owner not in clients:
            clients[owner] = {"insured": row[3], "inforce": 0, "plans": []}
        if row[8] == "Inforce":
            clients[owner]["inforce"] += 1
            clients[owner]["plans"].append(row[5])
    data = [
        [owner, info["insured"], str(info["inforce"]), "; ".join(info["plans"])]
        for owner, info in sorted(clients.items())
    ]
    nrows = write_table(ws, headers, data)
    style_sheet(ws, len(headers), nrows)


def build_readme(wb: Workbook) -> None:
    ws = wb.create_sheet("How to Use", 0)
    lines = [
        ("Client Base — Policy Book (Screenshots Only)", ""),
        ("", ""),
        ("Data source:", "Your policy book screenshots only"),
        ("", "No prospect/network list included"),
        ("", ""),
        ("Policies", "Same columns as your screenshots — one row per policy"),
        ("Clients Summary", "Auto-grouped from Policies (inforce count + plans)"),
        ("", ""),
        ("Regenerate:", "python3 build_excel.py"),
    ]
    for i, (a, b) in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=a)
        ws.cell(row=i, column=2, value=b)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 50
    ws["A1"].font = TITLE_FONT


def main() -> None:
    wb = Workbook()
    build_readme(wb)
    build_policies(wb)
    build_clients_summary(wb)
    wb.save(OUTPUT)
    print(f"Created {OUTPUT} (screenshot data only)")


if __name__ == "__main__":
    main()
